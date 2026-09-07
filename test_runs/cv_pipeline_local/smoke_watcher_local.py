import os
import shutil
from pathlib import Path


ROOT = Path.cwd() / "test_runs" / "cv_pipeline_local" / "files"
STAGING = ROOT / "staging"
LIBRARY = ROOT / "CV_Theque"
CV_DIR = STAGING / "data_analyst" / "junior"


def reset_sandbox() -> None:
    if ROOT.exists():
        shutil.rmtree(ROOT)
    CV_DIR.mkdir(parents=True, exist_ok=True)
    (STAGING / "processed").mkdir(parents=True, exist_ok=True)
    (STAGING / "failed").mkdir(parents=True, exist_ok=True)
    LIBRARY.mkdir(parents=True, exist_ok=True)


def main() -> None:
    reset_sandbox()

    os.environ["CV_STORAGE_ROOT"] = str(ROOT)
    os.environ["CV_LIBRARY_ROOT"] = str(LIBRARY)
    os.environ["WATCHER_STAGING_PATH"] = str(STAGING)
    os.environ["OPENROUTER_API_KEY"] = "test-only"

    from docx import Document

    sample = CV_DIR / "ali_smoke_cv.docx"
    doc = Document()
    doc.add_paragraph("Ali Smoke Test")
    doc.add_paragraph(
        "Data analyst junior profile with SQL Python Power BI dashboards and data quality testing. " * 4
    )
    doc.save(sample)

    import script.staging_watcher as watcher
    from service.cv_storage import LocalCVStorage

    storage = LocalCVStorage(ROOT, LIBRARY, STAGING)
    scanner = watcher.StagingScanner(storage, str(STAGING))
    items = scanner.scan()
    print("SCANNED", len(items), items)
    assert len(items) == 1, items
    assert items[0].profile_hint == "Data_Analyst", items[0]
    assert items[0].seniority_hint == "Junior", items[0]

    payload = {
        "informations_personnelles": {
            "nom_complet": "Ali Smoke Test",
            "titre": "Data Analyst",
        },
        "profil_resume": {"annees_experience": "1 an"},
        "competences": {"techniques": ["SQL", "Python", "Power BI"]},
    }

    watcher._extract_text_docx = lambda path: "Data analyst junior SQL Python Power BI reporting. " * 6
    watcher._extract_cv_json = lambda text: dict(payload)
    watcher._validate_json = lambda data: data
    watcher._classify = (
        lambda extracted, profile_hint, seniority_hint, known_profiles, pending_profiles, pending_lock:
        (profile_hint, seniority_hint)
    )

    processor = watcher.CVProcessor(storage)
    processor.process(items[0], frozenset({"Data_Analyst"}), {}, watcher.threading.Lock())

    processed = STAGING / "processed" / "data_analyst" / "junior" / "ali_smoke_cv.docx"
    failed = STAGING / "failed" / "data_analyst" / "junior" / "ali_smoke_cv.docx"
    json_path = LIBRARY / "Data_Analyst" / "Junior" / "extracted" / "ali_smoke_cv.json"
    original_path = LIBRARY / "Data_Analyst" / "Junior" / "originals" / "ali_smoke_cv.docx"

    print("PROCESSED_EXISTS", processed.exists(), processed)
    print("FAILED_EXISTS", failed.exists(), failed)
    print("JSON_EXISTS", json_path.exists(), json_path)
    print("ORIGINAL_EXISTS", original_path.exists(), original_path)

    assert processed.exists(), processed
    assert not failed.exists(), failed
    assert json_path.exists(), json_path
    assert original_path.exists(), original_path

    bad_dir = STAGING / "frontend" / "senior"
    bad_dir.mkdir(parents=True, exist_ok=True)
    bad_sample = bad_dir / "bad_cv.docx"
    bad_doc = Document()
    bad_doc.add_paragraph("Too short")
    bad_doc.save(bad_sample)

    bad_items = scanner.scan()
    print("SCANNED_AFTER_SUCCESS", len(bad_items), bad_items)
    assert len(bad_items) == 1, bad_items
    assert bad_items[0].profile_hint == "Frontend", bad_items[0]
    assert bad_items[0].seniority_hint == "Senior", bad_items[0]

    watcher._extract_text_docx = lambda path: "short"
    processor.process(bad_items[0], frozenset({"Frontend"}), {}, watcher.threading.Lock())

    bad_processed = STAGING / "processed" / "frontend" / "senior" / "bad_cv.docx"
    bad_failed = STAGING / "failed" / "frontend" / "senior" / "bad_cv.docx"
    print("BAD_PROCESSED_EXISTS", bad_processed.exists(), bad_processed)
    print("BAD_FAILED_EXISTS", bad_failed.exists(), bad_failed)

    assert not bad_processed.exists(), bad_processed
    assert bad_failed.exists(), bad_failed
    print("LOCAL_WATCHER_SMOKE_OK")


if __name__ == "__main__":
    main()
