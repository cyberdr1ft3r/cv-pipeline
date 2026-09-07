# script/agentic_extractor.py

"""
Agentic LLM-Based Text Extraction Module

Provides a unified, LLM-driven extraction pipeline for text files, images, and Excel files.
Prompts are loaded from external files defined in CONFIG["prompts"].
"""

import os
import yaml
import logging
import base64
import time
from abc import ABC, abstractmethod
from pathlib import Path
from typing import Optional, Tuple, Dict, Any, List
from dataclasses import dataclass, field
from enum import Enum
from openai import OpenAI
from dotenv import load_dotenv

# Configuration
ENV_PATH = Path(__file__).resolve().parents[1] / "config" / ".env"
load_dotenv(dotenv_path=ENV_PATH)

OPENROUTER_API_KEY = os.getenv("OPENROUTER_API_KEY")
if not OPENROUTER_API_KEY:
    raise RuntimeError(f"OPENROUTER_API_KEY is not defined. Expected in {ENV_PATH}")

PROJECT_ROOT = Path(__file__).resolve().parents[1]
CONFIG_PATH = PROJECT_ROOT / "config" / "config_yaml.yaml"

with open(CONFIG_PATH, "r", encoding="utf-8") as f:
    CONFIG = yaml.safe_load(f)

client = OpenAI(base_url=CONFIG["api"]["base_url"], api_key=OPENROUTER_API_KEY)


class FileType(Enum):
    TEXT = "text"
    IMAGE = "image"
    EXCEL = "excel"
    UNKNOWN = "unknown"


class ExtractionMethod(Enum):
    LLM_TEXT = "llm_text"
    LLM_VISION = "llm_vision"
    LLM_EXCEL = "llm_excel"
    FALLBACK_TEXT = "fallback_text"
    FALLBACK_OCR = "fallback_ocr"
    FALLBACK_EXCEL = "fallback_excel"
    FAILED = "failed"


@dataclass
class ExtractionResult:
    success: bool
    text: str
    method: ExtractionMethod
    file_type: FileType
    model_used: Optional[str] = None
    extraction_time_ms: int = 0
    metadata: Dict[str, Any] = field(default_factory=dict)
    
    def to_tuple(self) -> Tuple[bool, str, str]:
        return (self.success, self.text, self.method.value)


def setup_logger() -> logging.Logger:
    logger = logging.getLogger("AgenticExtractor")
    logger.setLevel(logging.DEBUG)
    logger.handlers = []
    console_handler = logging.StreamHandler()
    console_handler.setLevel(logging.INFO)
    console_handler.setFormatter(logging.Formatter("%(message)s"))
    logger.addHandler(console_handler)
    return logger

logger = setup_logger()


def load_prompt(prompt_key: str) -> str:
    """Load prompt from external file defined in CONFIG["prompts"]."""
    prompt_file = CONFIG["prompts"].get(prompt_key)
    if not prompt_file:
        raise ValueError(f"Prompt key '{prompt_key}' not found in config")
    
    prompt_path = PROJECT_ROOT / "config" / prompt_file
    if not prompt_path.exists():
        raise FileNotFoundError(f"Prompt file not found: {prompt_path}")
    
    with open(prompt_path, "r", encoding="utf-8") as f:
        return f.read()


VISION_MODELS = [
    "nvidia/nemotron-3-nano-30b-a3b:free",
    "google/gemini-2.0-flash-exp:free",
    "nvidia/nemotron-nano-12b-v2-vl:free",
]


class BaseExtractionAgent(ABC):
    def __init__(self, config: Dict[str, Any]):
        self.config = config
        self.primary_model = config["api"]["model"]
        self.fallback_models = config["api"].get("fallback_models", [])
        self.max_retries = config["api"].get("max_retries", 3)
        self.retry_wait = config["api"].get("retry_wait_seconds", 5)
        self.temperature = config["api"].get("temperature", 0.0)
        self.max_tokens = config["api"].get("max_tokens", 8192)
    
    @abstractmethod
    def extract(self, file_path: Path) -> ExtractionResult:
        pass
    
    @abstractmethod
    def get_file_type(self) -> FileType:
        pass
    
    def _call_llm(self, messages: List[Dict], model: Optional[str] = None) -> Tuple[bool, str, str]:
        models_to_try = [model or self.primary_model] + self.fallback_models
        last_error = None
        
        for current_model in models_to_try:
            for attempt in range(self.max_retries):
                try:
                    response = client.chat.completions.create(
                        model=current_model,
                        messages=messages,
                        max_tokens=self.max_tokens,
                        temperature=self.temperature
                    )
                    text = response.choices[0].message.content
                    if text and text.strip():
                        return True, text.strip(), current_model
                except Exception as e:
                    last_error = str(e)
                    logger.debug(f"   Model {current_model} attempt {attempt + 1} failed: {e}")
                    if attempt < self.max_retries - 1:
                        time.sleep(self.retry_wait)
        
        return False, f"All models failed. Last error: {last_error}", ""


class TextExtractionAgent(BaseExtractionAgent):
    def get_file_type(self) -> FileType:
        return FileType.TEXT
    
    def extract(self, file_path: Path) -> ExtractionResult:
        start_time = time.time()
        
        raw_content = self._read_file(file_path)
        if raw_content is None:
            return ExtractionResult(
                success=False,
                text=f"Failed to read file: {file_path}",
                method=ExtractionMethod.FAILED,
                file_type=self.get_file_type()
            )
        
        if len(raw_content) < 100 and raw_content.isascii():
            return ExtractionResult(
                success=True,
                text=raw_content,
                method=ExtractionMethod.FALLBACK_TEXT,
                file_type=self.get_file_type(),
                extraction_time_ms=int((time.time() - start_time) * 1000)
            )
        
        logger.info("   🤖 Cleaning text with LLM agent...")
        prompt_template = load_prompt("agentic_text")
        prompt = prompt_template.format(content=raw_content[:50000])
        
        messages = [{"role": "user", "content": prompt}]
        success, result, model = self._call_llm(messages)
        
        if success:
            return ExtractionResult(
                success=True,
                text=result,
                method=ExtractionMethod.LLM_TEXT,
                file_type=self.get_file_type(),
                model_used=model,
                extraction_time_ms=int((time.time() - start_time) * 1000)
            )
        
        logger.info("   ⚠️  LLM failed, using raw content as fallback...")
        return ExtractionResult(
            success=True,
            text=raw_content,
            method=ExtractionMethod.FALLBACK_TEXT,
            file_type=self.get_file_type(),
            extraction_time_ms=int((time.time() - start_time) * 1000),
            metadata={"llm_error": result}
        )
    
    def _read_file(self, file_path: Path) -> Optional[str]:
        encodings = ['utf-8', 'utf-16', 'latin-1', 'cp1252']
        for encoding in encodings:
            try:
                with open(file_path, 'r', encoding=encoding) as f:
                    return f.read().strip()
            except UnicodeDecodeError:
                continue
            except Exception:
                continue
        return None


class ImageExtractionAgent(BaseExtractionAgent):
    def get_file_type(self) -> FileType:
        return FileType.IMAGE
    
    def extract(self, file_path: Path) -> ExtractionResult:
        start_time = time.time()
        
        try:
            with open(file_path, 'rb') as f:
                image_data = base64.b64encode(f.read()).decode('utf-8')
        except Exception as e:
            return ExtractionResult(
                success=False,
                text=f"Failed to read image: {e}",
                method=ExtractionMethod.FAILED,
                file_type=self.get_file_type()
            )
        
        mime_type = self._get_mime_type(file_path)
        
        logger.info("   🤖 Extracting text with Vision LLM agent...")
        prompt = load_prompt("agentic_image")
        
        messages = [{
            "role": "user",
            "content": [
                {"type": "text", "text": prompt},
                {
                    "type": "image_url",
                    "image_url": {"url": f"data:{mime_type};base64,{image_data}"}
                }
            ]
        }]
        
        success, result, model = self._call_vision_llm(messages)
        
        if success and "NO_TEXT_FOUND" not in result:
            return ExtractionResult(
                success=True,
                text=result,
                method=ExtractionMethod.LLM_VISION,
                file_type=self.get_file_type(),
                model_used=model,
                extraction_time_ms=int((time.time() - start_time) * 1000)
            )
        
        logger.info("   ⚠️  Vision LLM failed, trying OCR fallback...")
        ocr_success, ocr_result = self._fallback_ocr(file_path)
        
        return ExtractionResult(
            success=ocr_success,
            text=ocr_result,
            method=ExtractionMethod.FALLBACK_OCR if ocr_success else ExtractionMethod.FAILED,
            file_type=self.get_file_type(),
            extraction_time_ms=int((time.time() - start_time) * 1000),
            metadata={"vision_error": result} if not success else {}
        )
    
    def _call_vision_llm(self, messages: List[Dict]) -> Tuple[bool, str, str]:
        last_error = None
        for model in VISION_MODELS:
            for attempt in range(2):
                try:
                    response = client.chat.completions.create(
                        model=model,
                        messages=messages,
                        max_tokens=CONFIG["api"]["vision"]["max_tokens"],
                        temperature=CONFIG["api"]["vision"]["temperature"],
                    )
                    text = response.choices[0].message.content
                    if text and text.strip() and len(text.strip()) > 20:
                        return True, text.strip(), model
                except Exception as e:
                    last_error = str(e)
                    logger.debug(f"   Vision model {model} failed: {e}")
                    time.sleep(2)
        return False, f"All vision models failed. Last error: {last_error}", ""
    
    def _get_mime_type(self, file_path: Path) -> str:
        mime_map = {
            '.png': 'image/png',
            '.jpg': 'image/jpeg',
            '.jpeg': 'image/jpeg',
            '.webp': 'image/webp',
            '.gif': 'image/gif',
            '.bmp': 'image/bmp',
            '.tiff': 'image/tiff',
        }
        return mime_map.get(file_path.suffix.lower(), 'image/png')
    
    def _fallback_ocr(self, file_path: Path) -> Tuple[bool, str]:
        try:
            import pytesseract
            from PIL import Image
            
            _tesseract_cmd = CONFIG.get("paths", {}).get("tesseract_cmd") or "tesseract"
            pytesseract.pytesseract.tesseract_cmd = _tesseract_cmd

            with Image.open(file_path) as image:
                if image.mode in ('RGBA', 'LA', 'P'):
                    background = Image.new('RGB', image.size, (255, 255, 255))
                    if image.mode == 'P':
                        image = image.convert('RGBA')
                    background.paste(image, mask=image.split()[-1] if 'A' in image.mode else None)
                    image = background
                elif image.mode != 'RGB':
                    image = image.convert('RGB')
                
                text = pytesseract.image_to_string(image, lang='eng+fra')
                if text and text.strip():
                    return True, text.strip()
        except Exception as e:
            logger.debug(f"   OCR fallback failed: {e}")
        return False, "OCR extraction failed"


class ExcelExtractionAgent(BaseExtractionAgent):
    def get_file_type(self) -> FileType:
        return FileType.EXCEL
    
    def extract(self, file_path: Path) -> ExtractionResult:
        start_time = time.time()

        raw_data = self._read_excel_structure(file_path)
        if raw_data is None:
            logger.info("   ⚠️  Excel structure read failed, trying library fallback...")
            fallback_success, fallback_result = self._fallback_library(file_path)
            return ExtractionResult(
                success=fallback_success,
                text=fallback_result,
                method=ExtractionMethod.FALLBACK_EXCEL if fallback_success else ExtractionMethod.FAILED,
                file_type=self.get_file_type(),
                extraction_time_ms=int((time.time() - start_time) * 1000),
            )
        
        logger.info("   🤖 Processing Excel with LLM agent...")
        prompt_template = load_prompt("agentic_excel")
        prompt = prompt_template.format(content=raw_data[:50000])
        
        messages = [{"role": "user", "content": prompt}]
        success, result, model = self._call_llm(messages)
        
        if success:
            return ExtractionResult(
                success=True,
                text=result,
                method=ExtractionMethod.LLM_EXCEL,
                file_type=self.get_file_type(),
                model_used=model,
                extraction_time_ms=int((time.time() - start_time) * 1000)
            )
        
        logger.info("   ⚠️  LLM failed, using library fallback...")
        fallback_success, fallback_result = self._fallback_library(file_path)
        
        return ExtractionResult(
            success=fallback_success,
            text=fallback_result,
            method=ExtractionMethod.FALLBACK_EXCEL if fallback_success else ExtractionMethod.FAILED,
            file_type=self.get_file_type(),
            extraction_time_ms=int((time.time() - start_time) * 1000),
            metadata={"llm_error": result} if not success else {}
        )
    
    def _read_excel_structure(self, file_path: Path) -> Optional[str]:
        file_ext = file_path.suffix.lower()
        try:
            if file_ext in ['.xlsx', '.xlsm']:
                return self._read_xlsx_structure(file_path)
            elif file_ext == '.xls':
                return self._read_xls_structure(file_path)
            return None
        except Exception as e:
            logger.debug(f"   Failed to read Excel structure: {e}")
            return None
    
    def _read_xlsx_structure(self, file_path: Path) -> Optional[str]:
        try:
            import openpyxl
        except ImportError:
            return None
        
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            output_parts = []
            
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                sheet_data = []
                
                if len(wb.sheetnames) > 1:
                    sheet_data.append(f"\n=== SHEET: {sheet_name} ===\n")
                
                for row in sheet.iter_rows():
                    row_values = []
                    for cell in row:
                        if cell.value is not None:
                            row_values.append(f"[{cell.coordinate}]: {cell.value}")
                    if row_values:
                        sheet_data.append(" | ".join(row_values))
                
                output_parts.extend(sheet_data)
            
            wb.close()
            return "\n".join(output_parts)
        except Exception:
            return None
    
    def _read_xls_structure(self, file_path: Path) -> Optional[str]:
        try:
            import xlrd
        except ImportError:
            return None
        
        try:
            workbook = xlrd.open_workbook(file_path)
            output_parts = []
            
            for sheet_idx in range(workbook.nsheets):
                sheet = workbook.sheet_by_index(sheet_idx)
                sheet_data = []
                
                if workbook.nsheets > 1:
                    sheet_data.append(f"\n=== SHEET: {sheet.name} ===\n")
                
                for row_idx in range(sheet.nrows):
                    row_values = []
                    for col_idx in range(sheet.ncols):
                        cell_value = sheet.cell_value(row_idx, col_idx)
                        if cell_value:
                            col_letter = chr(65 + col_idx) if col_idx < 26 else f"C{col_idx}"
                            row_values.append(f"[{col_letter}{row_idx + 1}]: {cell_value}")
                    if row_values:
                        sheet_data.append(" | ".join(row_values))
                
                output_parts.extend(sheet_data)
            
            return "\n".join(output_parts)
        except Exception:
            return None
    
    def _fallback_library(self, file_path: Path) -> Tuple[bool, str]:
        try:
            try:
                from script.offer_extractor import extract_from_excel_file
            except ImportError:
                from offer_extractor import extract_from_excel_file
            return extract_from_excel_file(file_path)
        except Exception as e:
            return False, f"Library extraction failed: {e}"


class AgenticExtractor:
    """Unified agentic extraction pipeline for text, images, and Excel files."""
    
    def __init__(self, config: Optional[Dict[str, Any]] = None):
        self.config = config or CONFIG
        self.text_agent = TextExtractionAgent(self.config)
        self.image_agent = ImageExtractionAgent(self.config)
        self.excel_agent = ExcelExtractionAgent(self.config)
    
    def extract(self, file_path: Path) -> ExtractionResult:
        if not file_path.exists():
            return ExtractionResult(
                success=False,
                text=f"File not found: {file_path}",
                method=ExtractionMethod.FAILED,
                file_type=FileType.UNKNOWN
            )
        
        file_type = self._detect_file_type(file_path)
        logger.info(f"   📄 Detected file type: {file_type.value}")
        
        if file_type == FileType.TEXT:
            return self.text_agent.extract(file_path)
        elif file_type == FileType.IMAGE:
            return self.image_agent.extract(file_path)
        elif file_type == FileType.EXCEL:
            return self.excel_agent.extract(file_path)
        else:
            return ExtractionResult(
                success=False,
                text=f"Unsupported file type: {file_path.suffix}",
                method=ExtractionMethod.FAILED,
                file_type=FileType.UNKNOWN
            )
    
    def _detect_file_type(self, file_path: Path) -> FileType:
        suffix = file_path.suffix.lower()
        
        TEXT_EXTENSIONS = {'.txt', '.md', '.text'}
        IMAGE_EXTENSIONS = {'.png', '.jpg', '.jpeg', '.bmp', '.tiff', '.webp', '.gif'}
        EXCEL_EXTENSIONS = {'.xlsx', '.xls', '.xlsm'}
        
        if suffix in TEXT_EXTENSIONS:
            return FileType.TEXT
        elif suffix in IMAGE_EXTENSIONS:
            return FileType.IMAGE
        elif suffix in EXCEL_EXTENSIONS:
            return FileType.EXCEL
        else:
            try:
                with open(file_path, 'rb') as f:
                    header = f.read(8)
                    if header[:8] == b'\x89PNG\r\n\x1a\n':
                        return FileType.IMAGE
                    elif header[:2] == b'\xff\xd8':
                        return FileType.IMAGE
                    elif header[:6] in (b'GIF87a', b'GIF89a'):
                        return FileType.IMAGE
                
                with open(file_path, 'r', encoding='utf-8') as f:
                    f.read(1000)
                return FileType.TEXT
            except (UnicodeDecodeError, IOError):
                return FileType.UNKNOWN


def extract_with_agent(file_path: Path) -> Tuple[bool, str, str]:
    """Convenience function for agentic extraction."""
    extractor = AgenticExtractor()
    result = extractor.extract(file_path)
    return result.to_tuple()


if __name__ == "__main__":
    import sys
    
    print("\n" + "=" * 70)
    print("🤖 AGENTIC TEXT EXTRACTOR")
    print("=" * 70)
    
    if len(sys.argv) > 1:
        test_file = Path(sys.argv[1])
    else:
        test_file = Path(input("\n📄 Enter path to test file: ").strip().strip('"'))
    
    if not test_file.exists():
        print(f"\n❌ File not found: {test_file}")
        sys.exit(1)
    
    print(f"\n⏳ Extracting from: {test_file.name}")
    print("-" * 70)
    
    extractor = AgenticExtractor()
    result = extractor.extract(test_file)
    
    print(f"\n{'✅' if result.success else '❌'} Extraction {'succeeded' if result.success else 'failed'}")
    print(f"   Method: {result.method.value}")
    print(f"   File type: {result.file_type.value}")
    if result.model_used:
        print(f"   Model: {result.model_used}")
    print(f"   Time: {result.extraction_time_ms}ms")
    print(f"   Characters: {len(result.text):,}")
    
    print("\n" + "-" * 70)
    print("PREVIEW (first 1000 characters):")
    print("-" * 70)
    print(result.text[:1000] + ("..." if len(result.text) > 1000 else ""))
    print("-" * 70)
