from __future__ import annotations

import tempfile
import threading
import unittest
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch
from xml.etree import ElementTree

from docx import Document

from script import staging_watcher as watcher
from service.cv_storage import LocalCVStorage


WORD_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WORD_BODY = f"{{{WORD_NAMESPACE}}}body"
WORD_PARAGRAPH = f"{{{WORD_NAMESPACE}}}p"
WORD_RUN = f"{{{WORD_NAMESPACE}}}r"
WORD_TEXT = f"{{{WORD_NAMESPACE}}}t"
WORD_DRAWING = f"{{{WORD_NAMESPACE}}}drawing"
WORD_TEXTBOX = f"{{{WORD_NAMESPACE}}}txbxContent"


def _llm_response(content: str, finish_reason: str) -> SimpleNamespace:
    choice = SimpleNamespace(
        message=SimpleNamespace(content=content),
        finish_reason=finish_reason,
    )
    return SimpleNamespace(choices=[choice])


def _add_textbox_to_docx(path: Path, marker: str) -> None:
    """Inject a minimal DrawingML textbox paragraph into a valid DOCX."""
    replacement = path.with_suffix(".replacement.docx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        replacement, "w"
    ) as destination:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "word/document.xml":
                root = ElementTree.fromstring(data)
                body = root.find(WORD_BODY)
                assert body is not None

                outer_paragraph = ElementTree.Element(WORD_PARAGRAPH)
                outer_run = ElementTree.SubElement(outer_paragraph, WORD_RUN)
                drawing = ElementTree.SubElement(outer_run, WORD_DRAWING)
                textbox = ElementTree.SubElement(drawing, WORD_TEXTBOX)
                textbox_paragraph = ElementTree.SubElement(textbox, WORD_PARAGRAPH)
                textbox_run = ElementTree.SubElement(textbox_paragraph, WORD_RUN)
                textbox_text = ElementTree.SubElement(textbox_run, WORD_TEXT)
                textbox_text.text = marker

                section_properties = body.find(
                    f"{{{WORD_NAMESPACE}}}sectPr"
                )
                insert_at = (
                    list(body).index(section_properties)
                    if section_properties is not None
                    else len(body)
                )
                body.insert(insert_at, outer_paragraph)
                data = ElementTree.tostring(
                    root,
                    encoding="utf-8",
                    xml_declaration=True,
                )
            destination.writestr(item, data)
    replacement.replace(path)


class DocxExtractionSafetyTests(unittest.TestCase):
    def test_extracts_body_textbox_header_and_footer_exactly_once(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "synthetic.docx"
            document = Document()
            document.add_paragraph("NORMAL-PARAGRAPH-MARKER")
            document.sections[0].header.paragraphs[0].text = "HEADER-MARKER"
            document.sections[0].footer.paragraphs[0].text = "FOOTER-MARKER"
            document.save(path)
            _add_textbox_to_docx(path, "TEXTBOX-MARKER")

            extracted = watcher._extract_text_docx(str(path))

        for marker in (
            "NORMAL-PARAGRAPH-MARKER",
            "TEXTBOX-MARKER",
            "HEADER-MARKER",
            "FOOTER-MARKER",
        ):
            with self.subTest(marker=marker):
                self.assertEqual(extracted.count(marker), 1)
        self.assertIn("NORMAL-PARAGRAPH-MARKER\n\nTEXTBOX-MARKER", extracted)

    def test_legacy_doc_is_never_opened_as_docx(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "legacy.doc"
            path.write_bytes(b"legacy-binary-document")

            with patch("docx.Document") as document:
                with self.assertRaisesRegex(
                    watcher.UnsupportedLegacyDocError,
                    r"unsupported legacy \.doc",
                ):
                    watcher._extract_text_docx(str(path))

            document.assert_not_called()


class LlmTruncationSafetyTests(unittest.TestCase):
    def setUp(self) -> None:
        self.api_config = patch.dict(
            watcher._CONFIG["api"],
            {"max_retries": 1, "retry_wait_seconds": 0},
        )
        self.api_config.start()

    def tearDown(self) -> None:
        self.api_config.stop()

    def test_stop_finish_reason_returns_normal_content(self) -> None:
        response = _llm_response('{"ok": true}', "stop")
        with patch.object(
            watcher._llm_client.chat.completions,
            "create",
            return_value=response,
        ):
            result = watcher._call_llm("safe prompt", reject_truncated=True)

        self.assertEqual(result, '{"ok": true}')

    def test_length_finish_reason_is_rejected_without_sensitive_logs(self) -> None:
        prompt_secret = "PROMPT-SECRET-9f81"
        response_secret = "RAW-RESPONSE-SECRET-83a2"
        token_secret = "TOKEN-SECRET-b7e4"
        response = _llm_response(response_secret, "length")

        with patch.object(watcher, "_OPENROUTER_API_KEY", token_secret):
            with patch.object(
                watcher._llm_client.chat.completions,
                "create",
                return_value=response,
            ):
                with self.assertLogs("staging_watcher", level="ERROR") as logs:
                    result = watcher._call_llm(
                        prompt_secret,
                        reject_truncated=True,
                    )

        self.assertIsNone(result)
        output = "\n".join(logs.output)
        self.assertIn(watcher._CONFIG["api"]["model"], output)
        self.assertIn("attempt=1/1", output)
        self.assertIn("finish_reason=length", output)
        self.assertNotIn(prompt_secret, output)
        self.assertNotIn(response_secret, output)
        self.assertNotIn(token_secret, output)

    def test_truncated_primary_never_reaches_parser_or_normalizer(self) -> None:
        response = _llm_response(
            '{"experiences_professionnelles": [], "projets_realises": [',
            "length",
        )
        with patch.object(
            watcher._llm_client.chat.completions,
            "create",
            return_value=response,
        ), patch.object(watcher, "_parse_llm_json") as parse_json, patch.object(
            watcher, "_validate_json"
        ) as validate_json:
            with self.assertRaisesRegex(RuntimeError, "no response"):
                extracted = watcher._extract_cv_json("CV text " * 30)
                validate_json(extracted)

        parse_json.assert_not_called()
        validate_json.assert_not_called()

    def test_truncated_validation_response_is_not_parsed(self) -> None:
        responses = [
            _llm_response("not valid JSON", "stop"),
            _llm_response('{"experiences_professionnelles": []', "length"),
        ]
        with patch.object(
            watcher._llm_client.chat.completions,
            "create",
            side_effect=responses,
        ), patch.object(
            watcher,
            "_parse_llm_json",
            side_effect=ValueError("invalid"),
        ) as parse_json:
            with self.assertRaisesRegex(ValueError, "Could not parse"):
                watcher._extract_cv_json("CV text " * 30)

        parse_json.assert_called_once_with("not valid JSON")

    def test_classification_keeps_default_non_rejecting_behavior(self) -> None:
        response = _llm_response("Security Engineer", "length")
        extracted = {"informations_personnelles": {"titre": "Cloud specialist"}}
        with patch.object(
            watcher._llm_client.chat.completions,
            "create",
            return_value=response,
        ):
            profile = watcher._resolve_profile(
                extracted,
                hint=None,
                known_profiles=frozenset(),
                pending_profiles={},
                pending_lock=threading.Lock(),
            )

        self.assertEqual(profile, "Security_Engineer")


class LegacyDocQuarantineTests(unittest.TestCase):
    def test_processor_preserves_legacy_doc_without_llm_or_candidate_data(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            staging = root / "staging"
            source = staging / "legacy-candidate.doc"
            source.parent.mkdir(parents=True)
            original = b"legacy-candidate-original"
            source.write_bytes(original)
            storage = LocalCVStorage(
                storage_root=root,
                library_root=root / "CV_Theque",
                staging_path=staging,
            )
            staged = watcher.StagedFile(
                str(source),
                source.name,
                None,
                None,
            )

            with patch.object(
                watcher,
                "WATCHER_STAGING_PATH",
                str(staging),
            ), patch.object(
                watcher,
                "_extract_text_docx",
            ) as extract_docx, patch.object(
                watcher,
                "_call_llm",
            ) as call_llm, patch.object(
                storage,
                "store_extracted",
                wraps=storage.store_extracted,
            ) as store_extracted, patch.object(
                storage,
                "store_original",
                wraps=storage.store_original,
            ) as store_original:
                with self.assertLogs("staging_watcher", level="ERROR") as logs:
                    watcher.CVProcessor(storage).process(
                        staged,
                        known_profiles=frozenset(),
                        pending_profiles={},
                        pending_lock=threading.Lock(),
                    )

            failed = staging / "failed" / source.name
            self.assertFalse(source.exists())
            self.assertEqual(failed.read_bytes(), original)
            self.assertIn(
                "unsupported legacy .doc format",
                "\n".join(logs.output),
            )
            extract_docx.assert_not_called()
            call_llm.assert_not_called()
            store_extracted.assert_not_called()
            store_original.assert_not_called()


if __name__ == "__main__":
    unittest.main()
