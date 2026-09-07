from __future__ import annotations

import html
import json
import os
import re
import tempfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Callable
from uuid import uuid4

from docx import Document
from openai import APIStatusError, AuthenticationError, OpenAI
from pypdf import PdfReader

from service.config import DATA_ROOT
from script.cv_generator import (
    LOGO_PATH,
    TEMPLATE_DIR,
    generate_certifications_section,
    generate_competences_section,
    generate_experiences_section,
    generate_formation_section,
    generate_langues_section,
    generate_pdf_from_html,
    generate_projets_section,
    get_logo_base64,
    html_escape,
)


ALLOWED_ALIGNMENT_EXTENSIONS = {".pdf", ".docx", ".txt"}
ALIGNMENTS_DIR = DATA_ROOT / "cv_alignments"
MAX_TEXT_CHARS = 45000
MAX_ALIGNMENT_OFFER_CHARS = 16000
MAX_ALIGNMENT_CV_CHARS = 22000
PROJECT_ROOT = Path(__file__).resolve().parents[1]


@dataclass
class AlignmentResult:
    alignment_id: str
    filename: str
    output_path: Path
    offer_chars: int
    cv_chars: int


def extract_document_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".txt":
        return path.read_text(encoding="utf-8", errors="ignore")
    if ext == ".pdf":
        return _extract_pdf_text(path)
    if ext == ".docx":
        return _extract_docx_text(path)
    raise ValueError(f"Format non supporté : {ext or '(aucun)'}")


def align_cv_to_offer(
    offer_path: Path,
    cv_path: Path,
    language: str = "fr",
    alignment_id: str | None = None,
    source_offer_name: str | None = None,
    source_cv_name: str | None = None,
    progress_callback: Callable[[int, str], None] | None = None,
) -> AlignmentResult:
    if progress_callback:
        progress_callback(20, "Extraction du texte de l'offre")
    offer_text = _compact_text(extract_document_text(offer_path))
    if progress_callback:
        progress_callback(35, "Extraction du texte du CV")
    cv_text = _compact_text(extract_document_text(cv_path))

    if len(offer_text) < 80:
        raise ValueError("Le texte de l'offre est trop court ou illisible.")
    if len(cv_text) < 120:
        raise ValueError("Le texte du CV est trop court ou illisible.")

    if progress_callback:
        progress_callback(50, "Préparation du prompt")
    aligned_cv = _call_alignment_model(
        offer_text=offer_text[:MAX_ALIGNMENT_OFFER_CHARS],
        cv_text=cv_text[: int(os.getenv("CV_ALIGNMENT_INPUT_CV_CHARS", str(MAX_ALIGNMENT_CV_CHARS)))],
        language=language,
    )
    if progress_callback:
        progress_callback(85, "Création du CV formaté")
    aligned_markdown = _render_formatted_cv_html(
        aligned_cv,
        offer_text=offer_text[:6000],
        template_name="classic",
    )

    alignment_id = alignment_id or uuid4().hex
    display_cv_name = source_cv_name or cv_path.name
    display_offer_name = source_offer_name or offer_path.name
    safe_stem = _safe_stem(Path(display_cv_name).stem) or "cv_aligne"
    filename = f"{safe_stem}_aligne.html"
    output_dir = ALIGNMENTS_DIR / alignment_id
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / filename
    output_path.write_text(
        _render_alignment_html(
            title=f"{cv_path.stem} - CV aligné",
            aligned_markdown=aligned_markdown,
            source_cv=display_cv_name,
            source_offer=display_offer_name,
        ),
        encoding="utf-8",
    )

    pdf_path = output_dir / f"{safe_stem}_aligne.pdf"
    if progress_callback:
        progress_callback(92, "Génération du fichier téléchargeable")
    pdf_ok, _pdf_error = generate_pdf_from_html(aligned_markdown, pdf_path)
    if pdf_ok and pdf_path.exists():
        filename = pdf_path.name
        output_path = pdf_path

    return AlignmentResult(
        alignment_id=alignment_id,
        filename=filename,
        output_path=output_path,
        offer_chars=len(offer_text),
        cv_chars=len(cv_text),
    )


def create_alignment_workspace(
    offer_filename: str,
    offer_content: bytes,
    cv_filename: str,
    cv_content: bytes,
    user_id: str | None = None,
) -> tuple[str, Path, Path]:
    alignment_id = uuid4().hex
    output_dir = ALIGNMENTS_DIR / alignment_id
    input_dir = output_dir / "inputs"
    input_dir.mkdir(parents=True, exist_ok=True)
    offer_path = _write_alignment_input(input_dir, offer_filename, offer_content)
    cv_path = _write_alignment_input(input_dir, cv_filename, cv_content)
    write_alignment_status(
        alignment_id,
        {
            "status": "queued",
            "progress": 5,
            "stage_label": "En attente",
            "alignment_id": alignment_id,
            "user_id": user_id,
            "offer_filename": offer_filename,
            "cv_filename": cv_filename,
            "created_at": datetime.utcnow().isoformat(),
        },
    )
    return alignment_id, offer_path, cv_path


def write_alignment_status(alignment_id: str, status: dict) -> None:
    if not re.fullmatch(r"[a-f0-9]{32}", alignment_id or ""):
        raise ValueError("Identifiant invalide")
    output_dir = ALIGNMENTS_DIR / alignment_id
    output_dir.mkdir(parents=True, exist_ok=True)
    status["updated_at"] = datetime.utcnow().isoformat()
    (output_dir / "status.json").write_text(
        json.dumps(status, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def read_alignment_status(alignment_id: str) -> dict:
    if not re.fullmatch(r"[a-f0-9]{32}", alignment_id or ""):
        raise ValueError("Identifiant invalide")
    status_path = ALIGNMENTS_DIR / alignment_id / "status.json"
    if not status_path.exists():
        raise FileNotFoundError("Alignement introuvable")
    return json.loads(status_path.read_text(encoding="utf-8"))


def list_recent_alignments(user_id: str | None, limit: int = 5, include_all: bool = False) -> list[dict]:
    limit = max(1, min(limit, 20))
    if not ALIGNMENTS_DIR.exists():
        return []

    rows: list[dict] = []
    for status_path in ALIGNMENTS_DIR.glob("*/status.json"):
        try:
            data = json.loads(status_path.read_text(encoding="utf-8"))
        except Exception:
            continue
        owner = data.get("user_id")
        if not include_all and owner != user_id:
            continue
        data.setdefault("alignment_id", status_path.parent.name)
        rows.append(data)

    rows.sort(key=lambda item: item.get("created_at") or item.get("updated_at") or "", reverse=True)
    return rows[:limit]


def resolve_alignment_download(alignment_id: str) -> Path:
    if not re.fullmatch(r"[a-f0-9]{32}", alignment_id or ""):
        raise ValueError("Identifiant invalide")
    output_dir = ALIGNMENTS_DIR / alignment_id
    matches = list(output_dir.glob("*.pdf")) or list(output_dir.glob("*.html"))
    if not matches:
        raise FileNotFoundError("CV aligné introuvable")
    return matches[0]


def _write_alignment_input(input_dir: Path, filename: str, content: bytes) -> Path:
    safe_name = Path(filename).name
    ext = Path(safe_name).suffix.lower()
    if ext not in ALLOWED_ALIGNMENT_EXTENSIONS:
        raise ValueError(
            f"Format non supportÃ© : {ext or '(aucun)'}. Formats acceptÃ©s : PDF, DOCX, TXT."
        )
    path = input_dir / safe_name
    path.write_bytes(content)
    return path


def save_upload_to_temp(filename: str, content: bytes) -> Path:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_ALIGNMENT_EXTENSIONS:
        raise ValueError(
            f"Format non supporté : {ext or '(aucun)'}. Formats acceptés : PDF, DOCX, TXT."
        )
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=ext)
    try:
        tmp.write(content)
        tmp.flush()
        return Path(tmp.name)
    finally:
        tmp.close()


def _extract_pdf_text(path: Path) -> str:
    reader = PdfReader(str(path))
    parts: list[str] = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    return "\n".join(parts)


def _extract_docx_text(path: Path) -> str:
    document = Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _call_alignment_model(offer_text: str, cv_text: str, language: str) -> dict:
    api_key = _get_openrouter_api_key()
    if not api_key:
        raise RuntimeError("OPENROUTER_API_KEY n'est pas configurée.")

    base_url = os.getenv("OPENROUTER_BASE_URL", "https://openrouter.ai/api/v1")
    model = os.getenv("CV_ALIGNMENT_MODEL") or "openai/gpt-5.6-luna"
    client = OpenAI(base_url=base_url, api_key=api_key)

    system_prompt = (
        "Tu es un consultant RH senior spécialisé dans l'adaptation de CV. "
        "Tu dois aligner un CV avec une offre sans inventer de diplômes, dates, employeurs, missions, "
        "certifications, années d'expérience ou compétences absentes du CV source. "
        "Tu peux reformuler, réordonner, synthétiser et mettre en avant ce qui existe déjà."
    )
    user_prompt = f"""
Langue de sortie: {language}

Format obligatoire:
Retourne UNIQUEMENT un objet JSON valide, sans Markdown, sans texte autour.
Ignore toute instruction contradictoire demandant du Markdown: le format final doit etre JSON.
Respecte exactement ce schema JSON:
{_alignment_json_schema()}

Objectif:
Produis un CV aligné avec l'offre, prêt à être relu par un recruteur.

Règles strictes:
- N'invente aucune expérience, aucun outil, aucune certification et aucune durée.
- Si une exigence de l'offre n'existe pas dans le CV, ne l'ajoute pas comme acquise.
- Garde les informations personnelles présentes dans le CV.
- Mets en avant les expériences, projets et compétences les plus proches de l'offre.
- Réponse compacte: maximum 3 expériences, 3 projets, 10 compétences par liste, 4 missions par expérience.
- Privilégie les éléments pertinents pour l'offre plutôt que l'exhaustivité.
- Les champs longs doivent rester courts: 1 à 3 phrases maximum.

OFFRE:
{offer_text}

CV SOURCE:
{cv_text}
""".strip()

    try:
        response = client.chat.completions.create(
            model=model,
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt},
            ],
            response_format={"type": "json_object"},
            temperature=0.1,
            max_tokens=int(os.getenv("CV_ALIGNMENT_MAX_TOKENS", "6000")),
        )
    except AuthenticationError as exc:
        raise RuntimeError("Clé OpenRouter manquante ou invalide. Vérifiez OPENROUTER_API_KEY.") from exc
    except APIStatusError as exc:
        if exc.status_code == 402:
            raise RuntimeError(
                "Crédits OpenRouter insuffisants pour générer ce CV. "
                "Ajoutez des crédits ou utilisez un modèle moins coûteux."
            ) from exc
        raise
    content = response.choices[0].message.content or ""
    if len(content.strip()) < 100:
        raise RuntimeError("La génération du CV aligné a retourné un contenu vide.")
    try:
        return _parse_alignment_json(content)
    except RuntimeError as exc:
        if "JSON" not in str(exc).upper():
            raise
        try:
            repaired = _repair_alignment_json(client=client, model=model, broken_content=content)
            return _parse_alignment_json(repaired)
        except Exception as repair_exc:
            raise RuntimeError("Le modèle n'a pas retourné un JSON valide exploitable.") from repair_exc


def _repair_alignment_json(client: OpenAI, model: str, broken_content: str) -> str:
    prompt = f"""
Répare ce contenu pour retourner UNIQUEMENT un objet JSON valide.
Ne résume pas. Ne commente pas. Ne mets aucun Markdown.
Si une valeur manque ou est coupée, utilise une chaîne vide ou une liste vide.
Respecte ce schéma:
{_alignment_json_schema()}

CONTENU À RÉPARER:
{broken_content[:24000]}
""".strip()
    response = client.chat.completions.create(
        model=model,
        messages=[
            {"role": "system", "content": "Tu répares strictement du JSON invalide en JSON valide."},
            {"role": "user", "content": prompt},
        ],
        response_format={"type": "json_object"},
        temperature=0,
        max_tokens=int(os.getenv("CV_ALIGNMENT_REPAIR_MAX_TOKENS", "6000")),
    )
    repaired = response.choices[0].message.content or ""
    if len(repaired.strip()) < 20:
        raise RuntimeError("La réparation JSON du CV aligné a retourné un contenu vide.")
    return repaired


def _alignment_json_schema() -> str:
    return json.dumps(
        {
            "informations_personnelles": {"nom_complet": "", "titre": ""},
            "profil_resume": {"description": "", "annees_experience": "", "specialisations": []},
            "formation": [{"institution": "", "lieu": "", "dates": "", "diplome": "", "specialisation": "", "mention": "", "details": ""}],
            "certifications": [{"nom": "", "organisme": "", "date": "", "validite": "", "niveau": ""}],
            "competences": {"methodologies_et_outils": [], "technologies": []},
            "projets_realises": [
                {
                    "nom": "",
                    "client_ou_contexte": "",
                    "periode": "",
                    "description": "",
                    "role": "",
                    "equipe": "",
                    "realisations": [],
                    "technologies": [],
                    "methodologies": [],
                    "source": "",
                    "experience_liee": "",
                }
            ],
            "langues": [{"nom": "", "niveau": "", "certifications": ""}],
            "experiences_professionnelles": [
                {
                    "titre_poste": "",
                    "entreprise": "",
                    "lieu": "",
                    "dates": "",
                    "contexte": "",
                    "missions": [],
                    "environnement": [],
                }
            ],
        },
        ensure_ascii=False,
    )


def _parse_alignment_json(content: str) -> dict:
    cleaned = content.strip()
    if cleaned.startswith("```"):
        cleaned = re.sub(r"^```(?:json)?", "", cleaned, flags=re.IGNORECASE).strip()
        cleaned = re.sub(r"```$", "", cleaned).strip()
    start = cleaned.find("{")
    if start > 0:
        cleaned = cleaned[start:]
    cleaned = _complete_partial_json(cleaned)
    cleaned = re.sub(r",\s*([}\]])", r"\1", cleaned)
    try:
        data = json.loads(cleaned)
    except json.JSONDecodeError as exc:
        raise RuntimeError("La generation du CV aligne n'a pas retourne un JSON valide.") from exc
    if not isinstance(data, dict):
        raise RuntimeError("La generation du CV aligne a retourne un format invalide.")
    return _normalize_alignment_json(data)


def _complete_partial_json(content: str) -> str:
    """Best-effort closure for provider responses truncated near valid JSON end."""
    cleaned = content.strip()
    if not cleaned:
        return cleaned

    in_string = False
    escaped = False
    stack: list[str] = []
    last_safe_index = -1

    for index, char in enumerate(cleaned):
        if escaped:
            escaped = False
            continue
        if char == "\\" and in_string:
            escaped = True
            continue
        if char == '"':
            in_string = not in_string
            continue
        if in_string:
            continue
        if char in "{[":
            stack.append("}" if char == "{" else "]")
        elif char in "}]":
            if stack and stack[-1] == char:
                stack.pop()
                last_safe_index = index
            else:
                break

    if not stack and last_safe_index >= 0:
        return cleaned[: last_safe_index + 1]
    if in_string:
        cleaned += '"'
    cleaned = re.sub(r",\s*$", "", cleaned)
    return cleaned + "".join(reversed(stack))


def _normalize_alignment_json(data: dict) -> dict:
    data.setdefault("informations_personnelles", {})
    data.setdefault("profil_resume", {})
    data.setdefault("formation", [])
    data.setdefault("certifications", [])
    data.setdefault("competences", {})
    data.setdefault("projets_realises", [])
    data.setdefault("langues", [])
    data.setdefault("experiences_professionnelles", [])

    info = data["informations_personnelles"]
    if not isinstance(info, dict):
        data["informations_personnelles"] = info = {}
    if not info.get("nom_complet"):
        info["nom_complet"] = "Candidat"

    competences = data["competences"]
    if not isinstance(competences, dict):
        data["competences"] = competences = {}
    competences.setdefault("methodologies_et_outils", [])
    competences.setdefault("technologies", [])

    for key in ["formation", "certifications", "projets_realises", "langues", "experiences_professionnelles"]:
        if not isinstance(data.get(key), list):
            data[key] = []
    for exp in data["experiences_professionnelles"]:
        if isinstance(exp, dict) and "environnement" not in exp:
            exp["environnement"] = exp.get("environnement_technique", [])
    return data


def _render_formatted_cv_html(cv_data: dict, offer_text: str, template_name: str = "classic") -> str:
    template_path = TEMPLATE_DIR / f"{template_name}.html"
    if not template_path.exists():
        template_path = TEMPLATE_DIR / "classic.html"
    template = template_path.read_text(encoding="utf-8")

    info = cv_data.get("informations_personnelles", {})
    nom_complet = html_escape(info.get("nom_complet", "CV"))
    titre = html_escape(info.get("titre", ""))
    logo_url = get_logo_base64(LOGO_PATH)
    left_content = "".join(
        [
            generate_formation_section(cv_data),
            generate_certifications_section(cv_data),
            generate_competences_section(cv_data),
            generate_langues_section(cv_data),
        ]
    )
    right_content_page1 = "".join(
        [
            _generate_static_profile_section(cv_data),
            generate_experiences_section(cv_data),
            generate_projets_section(cv_data),
        ]
    )
    return template.format(
        nom_complet=nom_complet,
        titre=titre,
        left_content=left_content,
        right_content_page1=right_content_page1,
        logo_url=logo_url,
        additional_pages="",
    )


def _generate_static_profile_section(data: dict) -> str:
    profil = data.get("profil_resume", {})
    if not isinstance(profil, dict):
        return ""
    description = html_escape(profil.get("description", ""))
    if not description:
        return ""
    return (
        '<div style="margin-bottom:0.4cm;padding:0.3cm;background:#f8f9fa;'
        'border-left:4px solid #3498db;border-radius:0 4px 4px 0">'
        '<h1 style="margin-top:0;margin-bottom:0.2cm">Profil</h1>'
        f'<div style="font-size:10pt;line-height:1.4;color:#2c3e50;text-align:justify">{description}</div>'
        "</div>"
    )


def _get_openrouter_api_key() -> str:
    value = os.getenv("OPENROUTER_API_KEY", "").strip()
    if value:
        return value

    env_path = PROJECT_ROOT / "config" / ".env"
    if not env_path.exists():
        return ""
    for line in env_path.read_text(encoding="utf-8-sig", errors="ignore").splitlines():
        line = line.strip().lstrip("\ufeff")
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, raw_value = line.split("=", 1)
        if key.strip() == "OPENROUTER_API_KEY":
            value = raw_value.strip().strip('"').strip("'")
            if value:
                os.environ["OPENROUTER_API_KEY"] = value
            return value
    return ""


def _render_alignment_html(title: str, aligned_markdown: str, source_cv: str, source_offer: str) -> str:
    if aligned_markdown.lstrip().lower().startswith(("<!doctype html", "<html")):
        return aligned_markdown

    body = _markdown_to_safe_html(aligned_markdown)
    generated_at = datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")
    return f"""<!doctype html>
<html lang="fr">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>{html.escape(title)}</title>
  <style>
    body {{ margin: 0; background: #f4f7fb; color: #0f172a; font-family: Arial, sans-serif; }}
    main {{ max-width: 900px; margin: 32px auto; background: white; padding: 42px; border: 1px solid #d8e0ea; }}
    h1 {{ font-size: 30px; margin: 0 0 16px; }}
    h2 {{ font-size: 20px; margin-top: 28px; border-bottom: 1px solid #e2e8f0; padding-bottom: 8px; }}
    h3 {{ font-size: 16px; margin-top: 20px; }}
    p, li {{ font-size: 14px; line-height: 1.65; }}
    ul {{ padding-left: 22px; }}
    .meta {{ margin-bottom: 28px; color: #64748b; font-size: 12px; }}
    strong {{ color: #0f172a; }}
    @media print {{ body {{ background: white; }} main {{ margin: 0; border: 0; }} }}
  </style>
</head>
<body>
  <main>
    <div class="meta">Source CV: {html.escape(source_cv)} · Offre: {html.escape(source_offer)} · Généré le {generated_at}</div>
    {body}
  </main>
</body>
</html>
"""


def _markdown_to_safe_html(markdown: str) -> str:
    lines = markdown.splitlines()
    out: list[str] = []
    in_list = False
    for raw in lines:
        line = raw.strip()
        if not line:
            if in_list:
                out.append("</ul>")
                in_list = False
            continue
        if line.startswith("### "):
            if in_list:
                out.append("</ul>")
                in_list = False
            out.append(f"<h3>{_inline_markdown(line[4:])}</h3>")
        elif line.startswith("## "):
            if in_list:
                out.append("</ul>")
                in_list = False
            out.append(f"<h2>{_inline_markdown(line[3:])}</h2>")
        elif line.startswith("# "):
            if in_list:
                out.append("</ul>")
                in_list = False
            out.append(f"<h1>{_inline_markdown(line[2:])}</h1>")
        elif line.startswith(("- ", "* ")):
            if not in_list:
                out.append("<ul>")
                in_list = True
            out.append(f"<li>{_inline_markdown(line[2:])}</li>")
        else:
            if in_list:
                out.append("</ul>")
                in_list = False
            out.append(f"<p>{_inline_markdown(line)}</p>")
    if in_list:
        out.append("</ul>")
    return "\n".join(out)


def _inline_markdown(value: str) -> str:
    escaped = html.escape(value)
    return re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", escaped)


def _compact_text(value: str) -> str:
    return re.sub(r"\n{3,}", "\n\n", value).strip()


def _safe_stem(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", value).strip("._-")[:80]
